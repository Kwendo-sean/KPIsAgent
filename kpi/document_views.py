"""Document Processing & Analysis views.

A separate module rather than more weight in views.py, which is already large.

The department workflow this serves — credit control chasing invoices, finance
filing remittances — needs the same guarantee as the rest of the app: amounts,
dates and aging are computed in Python by document_analysis. A language model
may describe a document, never value it.
"""
from __future__ import annotations

from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render

from .ai_agent import LocalAIUnavailable, local_ai_enabled
from .document_analysis import (
    build_aging_report,
    build_counterparty_exposure,
    extract_document_fields,
    summarize_queue,
)
from .models import DEPARTMENTS, DOCUMENT_STATUS, DOCUMENT_TYPES, Document
from .pdf_extractor import BankStatementPDFExtractor
from .views import _account_qs, _active_account, _audit, _base_context

_DOC_MAX_BYTES = 25 * 1024 * 1024

_VALID_TYPES = {code for code, _ in DOCUMENT_TYPES}
_VALID_DEPARTMENTS = {code for code, _ in DEPARTMENTS}
_VALID_STATUS = {code for code, _ in DOCUMENT_STATUS}


def _document_qs(request):
    """Documents scoped to the user and the active sub-account."""
    sub, _ = _active_account(request)
    qs = Document.objects.filter(owner=request.user)
    return qs.filter(sub_account=sub) if sub else qs.filter(sub_account__isnull=True)


def _apply_extracted_fields(document, text, file_name=""):
    """Fill a Document from deterministic extraction. No model involved."""
    fields = extract_document_fields(text or "", file_name)
    # Respect a type a person set; only fill in while still unclassified.
    if document.doc_type in ("", "OTHER") and fields["doc_type"]:
        document.doc_type = fields["doc_type"]
    document.amount = fields["amount"]
    document.currency = fields["currency"]
    document.doc_date = fields["doc_date"]
    document.due_date = fields["due_date"]
    document.counterparty = fields["counterparty"]
    document.reference = fields["reference"]
    return fields


def _extract_text(uploaded) -> tuple[str, str]:
    """Return (text, error). Reuses the statement extractor — no new deps."""
    name = (uploaded.name or "").lower()
    try:
        if name.endswith(".pdf"):
            uploaded.seek(0)
            text = BankStatementPDFExtractor.extract_text_from_pdf(uploaded) or ""
            if BankStatementPDFExtractor.needs_ocr(text):
                # Scanned page: local OCR only, and only when it is switched on.
                from .local_ocr import local_ocr_enabled, ocr_pages_to_text
                if local_ocr_enabled():
                    uploaded.seek(0)
                    images = BankStatementPDFExtractor.render_pdf_pages_to_images(uploaded)
                    ocr_text = ocr_pages_to_text(images) if images else None
                    if ocr_text:
                        return ocr_text, ""
                return text, (
                    "This PDF has little or no text layer. Enable local OCR to read "
                    "scanned documents, or upload a digital PDF."
                )
            return text, ""
        if name.endswith((".csv", ".txt")):
            uploaded.seek(0)
            return uploaded.read().decode("utf-8", errors="ignore"), ""
        if name.endswith((".xlsx", ".xls")):
            uploaded.seek(0)
            return _spreadsheet_to_text(uploaded, name), ""
    except Exception as exc:
        return "", f"Extraction failed: {exc}"[:400]

    return "", (
        "Text extraction is not supported for this file type. The document was "
        "stored, but no fields were extracted."
    )


def _spreadsheet_to_text(uploaded, name: str) -> str:
    """Flatten a spreadsheet so the same field parsers apply."""
    rows: list[str] = []
    if name.endswith(".xlsx"):
        import openpyxl
        workbook = openpyxl.load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(values_only=True):
            rows.append(",".join("" if c is None else str(c) for c in row))
    else:
        import xlrd
        book = xlrd.open_workbook(file_contents=uploaded.read())
        sheet = book.sheet_by_index(0)
        for idx in range(sheet.nrows):
            rows.append(",".join(str(sheet.cell_value(idx, c)) for c in range(sheet.ncols)))
    return "\n".join(rows)


@login_required(login_url="login")
def documents_view(request):
    """Document queue with credit-control aging and counterparty exposure."""
    documents = _document_qs(request).select_related("source_statement")

    department = request.GET.get("department", "").strip()
    doc_type = request.GET.get("doc_type", "").strip()
    status_filter = request.GET.get("status", "").strip()
    search = request.GET.get("q", "").strip()

    if department in _VALID_DEPARTMENTS:
        documents = documents.filter(department=department)
    if doc_type in _VALID_TYPES:
        documents = documents.filter(doc_type=doc_type)
    if status_filter in _VALID_STATUS:
        documents = documents.filter(status=status_filter)
    if search:
        documents = documents.filter(
            Q(title__icontains=search)
            | Q(counterparty__icontains=search)
            | Q(reference__icontains=search)
        )

    # Aging covers the whole open book, not the filtered view, so the headline
    # exposure does not shift as someone narrows the list.
    open_book = list(_document_qs(request))

    ctx = _base_context(request)
    ctx.update({
        "active_page": "documents",
        "documents": list(documents[:200]),
        "aging": build_aging_report(open_book),
        "exposure": build_counterparty_exposure(open_book),
        "queue": summarize_queue(open_book),
        "document_types": DOCUMENT_TYPES,
        "departments": DEPARTMENTS,
        "document_status_choices": DOCUMENT_STATUS,
        "filter_department": department,
        "filter_doc_type": doc_type,
        "filter_status": status_filter,
        "search_query": search,
        "importable_count": _account_qs(request).filter(documents__isnull=True).count(),
    })
    return render(request, "documents.html", ctx)


@login_required(login_url="login")
def import_statements_as_documents(request):
    """Create Document records from statements already in the system.

    Makes the section useful immediately without re-uploading anything.
    Statements that already have a Document are skipped, so running this twice
    is harmless.
    """
    if request.method != "POST":
        return redirect("documents")

    sub, _ = _active_account(request)
    pending = _account_qs(request).filter(documents__isnull=True)

    created = 0
    for statement in pending:
        document = Document(
            owner=request.user,
            sub_account=sub,
            title=statement.file_name or f"Statement {statement.pk}",
            doc_type="STATEMENT",
            department="FINANCE",
            status="PROCESSED" if statement.is_processed else "PENDING",
            source="STATEMENT",
            source_statement=statement,
            extracted_text=(statement.extracted_text or "")[:200000],
        )
        _apply_extracted_fields(document, statement.extracted_text or "",
                                statement.file_name or "")
        # The statement's own parsed figures beat regex over its text.
        if statement.total_withdrawals is not None:
            document.amount = statement.total_withdrawals
        if statement.statement_period_end:
            document.doc_date = statement.statement_period_end
        document.save()
        created += 1

    _audit(request, "DOCUMENT_IMPORT", f"Imported {created} statement(s) as documents")
    messages.success(
        request,
        f"Imported {created} statement(s) into Documents."
        if created else "No new statements to import — all are already listed.",
    )
    return redirect("documents")


@login_required(login_url="login")
def upload_document(request):
    """Upload a document, extract its text, then parse fields deterministically."""
    if request.method != "POST":
        return redirect("documents")

    uploaded = request.FILES.get("file")
    if not uploaded:
        messages.error(request, "Please choose a file to upload.")
        return redirect("documents")
    if uploaded.size > _DOC_MAX_BYTES:
        messages.error(request, "File is too large. Maximum size is 25 MB.")
        return redirect("documents")

    department = (request.POST.get("department") or "").strip()
    doc_type = (request.POST.get("doc_type") or "").strip()
    title = (request.POST.get("title") or "").strip()

    sub, _ = _active_account(request)
    document = Document(
        owner=request.user,
        sub_account=sub,
        title=title or uploaded.name,
        doc_type=doc_type if doc_type in _VALID_TYPES else "OTHER",
        department=department if department in _VALID_DEPARTMENTS else "CREDIT_CONTROL",
        source="UPLOAD",
        file=uploaded,
    )

    text, error = _extract_text(uploaded)
    document.extracted_text = (text or "")[:200000]
    document.processing_error = error
    _apply_extracted_fields(document, text, uploaded.name or "")
    document.status = "PENDING" if text else "FAILED"
    document.save()

    _audit(request, "DOCUMENT_UPLOAD", f"Uploaded document (id={document.pk})")
    if error:
        messages.warning(request, error)
    else:
        messages.success(request, f"Uploaded and read “{document.title}”.")
    return redirect("document_detail", document_id=document.pk)


@login_required(login_url="login")
def document_detail(request, document_id):
    """One document: extracted fields, review actions, optional AI summary."""
    document = get_object_or_404(Document, pk=document_id, owner=request.user)
    ctx = _base_context(request)
    ctx.update({
        "active_page": "documents",
        "document": document,
        "days_overdue": document.days_overdue(),
        "document_types": DOCUMENT_TYPES,
        "departments": DEPARTMENTS,
        "document_status_choices": DOCUMENT_STATUS,
        "text_preview": (document.extracted_text or "")[:4000],
        "local_ai": local_ai_enabled(),
    })
    return render(request, "document_detail.html", ctx)


@login_required(login_url="login")
def update_document(request, document_id):
    """Save a reviewer's corrections. A person always outranks the parser."""
    document = get_object_or_404(Document, pk=document_id, owner=request.user)
    if request.method != "POST":
        return redirect("document_detail", document_id=document.pk)

    status = (request.POST.get("status") or "").strip()
    department = (request.POST.get("department") or "").strip()
    doc_type = (request.POST.get("doc_type") or "").strip()
    if status in _VALID_STATUS:
        document.status = status
    if department in _VALID_DEPARTMENTS:
        document.department = department
    if doc_type in _VALID_TYPES:
        document.doc_type = doc_type

    document.counterparty = (request.POST.get("counterparty") or "").strip()[:200]
    document.reference = (request.POST.get("reference") or "").strip()[:100]
    document.notes = (request.POST.get("notes") or "").strip()[:5000]

    amount_raw = (request.POST.get("amount") or "").strip().replace(",", "")
    if amount_raw:
        try:
            document.amount = Decimal(amount_raw)
        except Exception:
            messages.error(request, "Amount was not a valid number and was left unchanged.")
    else:
        document.amount = None

    for field in ("doc_date", "due_date"):
        raw = (request.POST.get(field) or "").strip()
        setattr(document, field,
                BankStatementPDFExtractor._parse_date(raw) if raw else None)

    document.save()
    _audit(request, "DOCUMENT_UPDATE", f"Updated document {document.pk}")
    messages.success(request, "Document updated.")
    return redirect("document_detail", document_id=document.pk)


@login_required(login_url="login")
def reextract_document(request, document_id):
    """Re-run deterministic extraction over the stored text."""
    document = get_object_or_404(Document, pk=document_id, owner=request.user)
    if request.method != "POST":
        return redirect("document_detail", document_id=document.pk)

    _apply_extracted_fields(document, document.extracted_text or "", document.title)
    document.save()
    _audit(request, "DOCUMENT_REEXTRACT", f"Re-extracted document {document.pk}")
    messages.success(request, "Fields re-extracted from the stored text.")
    return redirect("document_detail", document_id=document.pk)


@login_required(login_url="login")
def summarize_document(request, document_id):
    """Have the model describe what was already extracted.

    Narrative only: it is shown the extracted fields, not the raw document, so
    it has nothing to compute from. Every figure on the page comes from the
    deterministic fields, and in local mode a model failure is surfaced rather
    than being allowed to escalate to a cloud provider.
    """
    document = get_object_or_404(Document, pk=document_id, owner=request.user)
    if request.method != "POST":
        return redirect("document_detail", document_id=document.pk)

    facts = {
        "type": document.get_doc_type_display(),
        "department": document.get_department_display(),
        "counterparty": document.counterparty or "not stated",
        "reference": document.reference or "not stated",
        "amount": (f"{document.currency or ''} {document.amount:,.2f}".strip()
                   if document.amount is not None else "not stated"),
        "document date": str(document.doc_date or "not stated"),
        "due date": str(document.due_date or "not stated"),
        "status": document.get_status_display(),
    }
    prompt = (
        "Summarise this document for a credit control officer in 2-3 sentences. "
        "Use ONLY the fields below. Do not add figures, dates, causes or "
        "recommendations that are not listed. Do not suggest fraud or wrongdoing.\n\n"
        + "\n".join(f"{k}: {v}" for k, v in facts.items())
    )

    try:
        from .ai_agent import _ask_text
        summary = _ask_text(prompt, 220)
        document.summary = (summary or "").strip()[:2000]
        document.save(update_fields=["summary", "updated_at"])
        messages.success(request, "Summary generated.") if document.summary else \
            messages.error(request, "The model returned no summary.")
    except LocalAIUnavailable as exc:
        messages.error(request, str(exc))
    except Exception as exc:
        messages.error(request, f"Summary failed: {exc}")

    return redirect("document_detail", document_id=document.pk)


@login_required(login_url="login")
def delete_document(request, document_id):
    document = get_object_or_404(Document, pk=document_id, owner=request.user)
    if request.method != "POST":
        return redirect("document_detail", document_id=document.pk)
    title = document.title
    document.delete()
    _audit(request, "DOCUMENT_DELETE", f"Deleted document {title}")
    messages.success(request, f"Deleted “{title}”.")
    return redirect("documents")
